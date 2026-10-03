from __future__ import annotations

import csv
import hashlib
import io
import json
import re
from pathlib import Path
from typing import Any

from app.config import CHUNK_OVERLAP, CHUNK_SIZE
from knowledge.ingest import chunk_text, normalize_whitespace

ALLOWED_SUFFIXES = {".txt", ".md", ".json", ".jsonl", ".csv", ".pdf", ".docx"}


def extract_text(filename: str, content: bytes) -> list[dict[str, str]]:
    """Extract text from supported files, preserving a useful title per record."""
    suffix = Path(filename).suffix.lower()
    if suffix not in ALLOWED_SUFFIXES:
        raise ValueError(f"Unsupported file type '{suffix or '(none)'}'. Supported: {', '.join(sorted(ALLOWED_SUFFIXES))}")
    if not content:
        raise ValueError("The uploaded file is empty.")
    title = Path(filename).stem
    if suffix in {".txt", ".md"}:
        return [{"title": title, "text": content.decode("utf-8-sig", errors="replace")}]
    if suffix == ".json":
        data = json.loads(content.decode("utf-8-sig"))
        data = data.get("documents", data) if isinstance(data, dict) else data
        data = data if isinstance(data, list) else [data]
        return _records(data, title)
    if suffix == ".jsonl":
        records = [json.loads(line) for line in content.decode("utf-8-sig").splitlines() if line.strip()]
        return _records(records, title)
    if suffix == ".csv":
        rows = list(csv.DictReader(io.StringIO(content.decode("utf-8-sig", errors="replace"))))
        records = []
        for row_number, row in enumerate(rows, 1):
            values = {str(key).strip(): str(value).strip() for key, value in row.items() if key and value not in (None, "")}
            if values:
                record_title = values.get("title") or values.get("name") or f"{title} — row {row_number}"
                body = values.get("text") or values.get("content") or "; ".join(f"{key}: {value}" for key, value in values.items())
                records.append({"title": record_title, "text": body})
        if not records:
            raise ValueError("CSV file contains no non-empty data rows.")
        return records
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(content))
        records = [{"title": f"{title} — page {i}", "text": page.extract_text() or ""} for i, page in enumerate(reader.pages, 1)]
        return [r for r in records if r["text"].strip()]
    from docx import Document
    doc = Document(io.BytesIO(content))
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        paragraphs.extend(" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows)
    return [{"title": title, "text": "\n".join(paragraphs)}]


def _records(data: list[Any], default_title: str) -> list[dict[str, str]]:
    records = []
    for item in data:
        if not isinstance(item, dict):
            continue
        text = item.get("text", item.get("content", item.get("description", "")))
        if isinstance(text, (dict, list)):
            text = json.dumps(text, ensure_ascii=False)
        if str(text).strip():
            records.append({"title": str(item.get("title") or default_title), "text": str(text)})
    if not records:
        raise ValueError("No usable text records were found in this file.")
    return records


def make_chunks(records: list[dict[str, str]], source: str) -> list[dict[str, Any]]:
    result = []
    for record_number, record in enumerate(records):
        clean = normalize_whitespace(record["text"])
        for chunk_number, chunk in enumerate(chunk_text(clean, CHUNK_SIZE, CHUNK_OVERLAP)):
            identity = hashlib.sha256(f"{source}\0{record_number}\0{chunk_number}\0{chunk}".encode()).hexdigest()[:20]
            result.append({
                "id": f"upload_{identity}", "title": record["title"], "text": chunk,
                "category": "general", "anomaly_type": "general", "source": source,
                "source_url": "", "source_file": source, "chunk_index": chunk_number,
                "chunk_size": len(chunk),
            })
    if not result:
        raise ValueError("No text could be extracted from the uploaded file.")
    return result
